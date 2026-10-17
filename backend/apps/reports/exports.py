"""Server-side export generation — Excel (openpyxl) and PDF (reportlab).

Exports run through the exact same scoping + filter pipeline as the on-screen
report; nothing here re-queries data with a wider scope.

Arabic correctness:
- Excel stores logical text — shaping is the spreadsheet app's job. Sheets are
  marked right-to-left when the report language is Arabic.
- PDF needs explicit shaping (contextual letter forms + bidi reorder): text is
  run through arabic_reshaper + python-bidi and rendered with the bundled Noto
  Sans Arabic / DejaVu fonts, split per-script so Latin stays in DejaVu.
"""
from html import escape
from io import BytesIO
from pathlib import Path

import arabic_reshaper
from bidi.algorithm import get_display
from django.utils import timezone
from rest_framework.exceptions import ValidationError

EXPORT_MAX_ROWS_XLSX = 50_000
EXPORT_MAX_ROWS_PDF = 2_000

CDS_RED = "#E31E24"
CDS_DEEP_RED = "#B7161B"
CDS_DARK_GRAY = "#4A4D52"

_FONT_DIR = Path(__file__).resolve().parent / "fonts"
_ARABIC_RE = None

TITLES = {
    "attendance": {"en": "Attendance Report", "ar": "تقرير الحضور"},
    "timesheets": {"en": "Timesheet Report", "ar": "تقرير كشوف الدوام"},
    "requests": {"en": "Requests Report", "ar": "تقرير الطلبات"},
    "employees": {"en": "Employee Summary Report", "ar": "تقرير ملخص الموظفين"},
}

FILTER_LABELS = {
    "date_from": {"en": "From", "ar": "من"},
    "date_to": {"en": "To", "ar": "إلى"},
    "employee": {"en": "Employee", "ar": "الموظف"},
    "department": {"en": "Department", "ar": "القسم"},
    "status": {"en": "Status", "ar": "الحالة"},
    "request_type": {"en": "Request Type", "ar": "نوع الطلب"},
}

# (row key, en header, ar header, type, pdf width weight)
COLUMNS = {
    "attendance": [
        ("employee_name", "Employee", "الموظف", "text", 3),
        ("department", "Department", "القسم", "text", 2),
        ("date", "Date", "التاريخ", "date", 2),
        ("check_in", "Check In", "تسجيل الدخول", "time", 2),
        ("check_out", "Check Out", "تسجيل الخروج", "time", 2),
        ("worked_hours", "Worked Hours", "ساعات العمل", "hours", 2),
        ("required_hours", "Required Hours", "الساعات المطلوبة", "hours", 2),
        ("difference_hours", "Difference", "الفرق", "hours", 2),
        ("status", "Status", "الحالة", "attendance_status", 2),
    ],
    "timesheets": [
        ("employee_name", "Employee", "الموظف", "text", 3),
        ("department", "Department", "القسم", "text", 2),
        ("week", "Week (Sun–Sat)", "الأسبوع (أحد–سبت)", "week", 3),
        ("logged_hours", "Logged Hours", "الساعات المسجلة", "hours", 2),
        ("required_hours", "Required Hours", "الساعات المطلوبة", "hours", 2),
        ("difference_hours", "Difference", "الفرق", "hours", 2),
        ("status", "Status", "الحالة", "item_status", 2),
        ("submitted_at", "Submitted At", "تاريخ التقديم", "datetime", 2),
        ("reviewed_at", "Reviewed At", "تاريخ المراجعة", "datetime", 2),
    ],
    "requests": [
        ("employee_name", "Employee", "الموظف", "text", 3),
        ("request_type", "Request Type", "نوع الطلب", "text", 3),
        ("created_at", "Created", "تاريخ الإنشاء", "datetime", 2),
        ("status", "Status", "الحالة", "item_status", 2),
        ("pending_with", "Pending With", "بانتظار", "actor_role", 2),
    ],
    "employees": [
        ("employee_name", "Employee", "الموظف", "text", 3),
        ("department", "Department", "القسم", "text", 2),
        ("attendance_days", "Attendance Days", "أيام الحضور", "int", 2),
        ("worked_hours", "Worked Hours", "ساعات العمل", "hours", 2),
        ("timesheets", "Timesheets", "كشوف الدوام", "int", 2),
        ("logged_hours", "Logged Hours", "الساعات المسجلة", "hours", 2),
        ("requests_total", "Requests", "الطلبات", "int", 2),
        ("requests_pending", "Pending", "في الانتظار", "int", 2),
        ("requests_approved", "Approved", "تمت الموافقة", "int", 2),
        ("requests_rejected", "Rejected", "مرفوض", "int", 2),
        ("requests_cancelled", "Cancelled", "تم الإلغاء", "int", 2),
    ],
}

ITEM_STATUS_LABELS = {
    "PENDING": {"en": "Pending", "ar": "في الانتظار"},
    "IN_REVIEW": {"en": "In Review", "ar": "قيد المراجعة"},
    "RETURNED": {"en": "Returned", "ar": "معاد للتصحيح"},
    "APPROVED": {"en": "Approved", "ar": "تمت الموافقة"},
    "REJECTED": {"en": "Rejected", "ar": "مرفوض"},
    "WITHDRAWN": {"en": "Cancelled", "ar": "تم الإلغاء"},
}

ATTENDANCE_STATUS_LABELS = {
    "PRESENT": {"en": "Present", "ar": "حاضر"},
    "IN_PROGRESS": {"en": "In Progress", "ar": "قيد العمل"},
}

ROLE_LABELS = {
    "MANAGER": {"en": "Manager", "ar": "المدير"},
    "HR": {"en": "HR", "ar": "الموارد البشرية"},
}

_SUMMARY_LABELS = {
    "total_worked_hours": {"en": "Total Worked Hours", "ar": "إجمالي ساعات العمل"},
    "required_hours": {"en": "Required Hours", "ar": "الساعات المطلوبة"},
    "attendance_days": {"en": "Attendance Days", "ar": "أيام الحضور"},
    "employees": {"en": "Employees", "ar": "الموظفون"},
    "sessions": {"en": "Sessions", "ar": "الجلسات"},
    "late_days": {"en": "Late Days", "ar": "أيام التأخير"},
    "overtime_hours": {"en": "Overtime Hours", "ar": "ساعات إضافية"},
    "timesheets": {"en": "Timesheets", "ar": "كشوف الدوام"},
    "total_logged_hours": {"en": "Total Logged Hours", "ar": "إجمالي الساعات المسجلة"},
    "total": {"en": "Total Requests", "ar": "إجمالي الطلبات"},
    "requests_total": {"en": "Requests", "ar": "الطلبات"},
    "requests_pending": {"en": "Requests Pending", "ar": "طلبات في الانتظار"},
    "requests_approved": {"en": "Requests Approved", "ar": "طلبات معتمدة"},
    "requests_rejected": {"en": "Requests Rejected", "ar": "طلبات مرفوضة"},
    "requests_cancelled": {"en": "Requests Cancelled", "ar": "طلبات ملغاة"},
}


def _has_arabic(text):
    global _ARABIC_RE
    if _ARABIC_RE is None:
        import re

        _ARABIC_RE = re.compile(
            r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]"
        )
    return bool(text) and bool(_ARABIC_RE.search(str(text)))


def shape_arabic(text):
    """Logical Arabic → display order (reshaped glyphs, bidi applied)."""
    return get_display(arabic_reshaper.reshape(str(text)))


def _fmt_hours(v):
    if v is None:
        return "—"
    return f"{float(v):.2f}"


def _cell_value(row, key, col_type, lang):
    v = row.get(key)
    if key == "week":
        start, end = row.get("business_week_start"), row.get("week_end")
        return f"{start} → {end}"
    if key == "pending_with":
        role = row.get("current_actor_role")
        if not role or row.get("status") not in ("PENDING", "IN_REVIEW"):
            return "—"
        return ROLE_LABELS.get(role, {}).get(lang, role)
    if v is None or v == "":
        return "—" if col_type in ("hours", "datetime", "time", "actor_role") else ""
    if col_type == "hours":
        return _fmt_hours(v)
    if col_type in ("date",):
        return str(v)[:10]
    if col_type == "time":
        return timezone.localtime(v).strftime("%H:%M") if hasattr(v, "strftime") else str(v)
    if col_type == "datetime":
        return (
            timezone.localtime(v).strftime("%Y-%m-%d %H:%M")
            if hasattr(v, "strftime")
            else str(v)
        )
    if col_type == "item_status":
        return ITEM_STATUS_LABELS.get(v, {}).get(lang, v)
    if col_type == "attendance_status":
        return ATTENDANCE_STATUS_LABELS.get(v, {}).get(lang, v)
    return str(v)


def _filters_text(applied, lang, employee_names=None, request_type_names=None):
    parts = []
    for key, value in applied.items():
        label = FILTER_LABELS.get(key, {}).get(lang, key)
        shown = value
        if key == "employee" and employee_names:
            shown = employee_names.get(value, str(value))
        if key == "request_type" and request_type_names:
            shown = request_type_names.get(value, str(value))
        parts.append(f"{label}: {shown}")
    return "  ·  ".join(parts) if parts else (
        "لا توجد" if lang == "ar" else "None"
    )


def _summary_pairs(report_key, summary, lang):
    pairs = []
    for key, value in summary.items():
        if key.endswith("_available") or key in ("by_status", "by_status_group",
                                                "by_request_type", "by_type_status"):
            continue
        label = _SUMMARY_LABELS.get(key, {}).get(lang, key.replace("_", " ").title())
        pairs.append((label, "—" if value is None else value))
    for key, bucket in (("by_status", "Status"), ("by_request_type", "Type")):
        data = summary.get(key) or {}
        for name, n in data.items():
            if key == "by_status":
                name = ITEM_STATUS_LABELS.get(name, {}).get(lang, name)
            pairs.append((name, n))
    return pairs


def check_export_size(report_key, fmt, row_count):
    limit = EXPORT_MAX_ROWS_PDF if fmt == "pdf" else EXPORT_MAX_ROWS_XLSX
    if row_count > limit:
        raise ValidationError(
            {"export": f"Result has {row_count} rows (limit {limit} for "
                       f"{fmt.upper()}). Narrow the filters and retry."}
        )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def build_xlsx(report_key, *, rows, summary, meta, lang):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rtl = lang == "ar"
    title = TITLES[report_key][lang]
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.rightToLeft = rtl

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=CDS_RED[1:])
    align_r = Alignment(horizontal="right" if rtl else "left")

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color=CDS_RED[1:])
    ws.cell(row=2, column=1, value=("تاريخ الإنشاء" if rtl else "Generated"))
    ws.cell(row=2, column=2, value=timezone.localtime().strftime("%Y-%m-%d %H:%M"))
    ws.cell(row=3, column=1, value=("الفلاتر" if rtl else "Filters"))
    ws.cell(row=3, column=2, value=_filters_text(meta.get("filters", {}), lang))

    cols = COLUMNS[report_key]
    header_row = 5
    for i, (key, en, ar, _t, _w) in enumerate(cols, start=1):
        c = ws.cell(row=header_row, column=i, value=ar if rtl else en)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = f"A{header_row + 1}"

    for r_i, row in enumerate(rows, start=header_row + 1):
        for c_i, (key, _en, _ar, col_type, _w) in enumerate(cols, start=1):
            c = ws.cell(
                row=r_i, column=c_i,
                value=_cell_value(row, key, col_type, lang),
            )
            c.alignment = align_r

    for c_i, (_k, en, ar, _t, w) in enumerate(cols, start=1):
        label = ar if rtl else en
        ws.column_dimensions[get_column_letter(c_i)].width = max(12, len(label) + 4, w * 6)

    ws2 = wb.create_sheet("Summary")
    ws2.sheet_view.rightToLeft = rtl
    ws2.cell(row=1, column=1, value=("الملخص" if rtl else "Summary")).font = Font(
        bold=True, size=12, color=CDS_RED[1:]
    )
    for i, (label, value) in enumerate(_summary_pairs(report_key, summary, lang), start=2):
        ws2.cell(row=i, column=1, value=str(label))
        ws2.cell(row=i, column=2, value=value)
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 18

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_LATIN = "CDSLatin"
_LATIN_B = "CDSLatin-Bold"
_AR = "CDSArabic"
_AR_B = "CDSArabic-Bold"
_fonts_ready = False


def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont(_LATIN, str(_FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont(_LATIN_B, str(_FONT_DIR / "DejaVuSans-Bold.ttf")))
    pdfmetrics.registerFont(TTFont(_AR, str(_FONT_DIR / "NotoSansArabic-Regular.ttf")))
    pdfmetrics.registerFont(TTFont(_AR_B, str(_FONT_DIR / "NotoSansArabic-Bold.ttf")))
    _fonts_ready = True


def _pdf_text(text):
    """Return display-order markup for a Paragraph: Arabic runs are shaped,
    reordered via bidi, and wrapped in the Arabic font; Latin stays DejaVu."""
    s = str(text)
    if not _has_arabic(s):
        return escape(s)
    display = shape_arabic(s)
    runs, buf, in_ar = [], [], _has_arabic(display[0])
    for ch in display:
        ch_ar = _has_arabic(ch)
        if ch_ar != in_ar:
            runs.append((buf, in_ar))
            buf, in_ar = [], ch_ar
        buf.append(ch)
    runs.append((buf, in_ar))
    parts = []
    for chars, ar in runs:
        seg = escape("".join(chars))
        parts.append(f'<font name="{_AR}">{seg}</font>' if ar else seg)
    return "".join(parts)


def _pdf_title(text):
    s = str(text)
    if not _has_arabic(s):
        return escape(s)
    return f'<font name="{_AR_B}">{escape(shape_arabic(s))}</font>'


def build_pdf(report_key, *, rows, summary, meta, lang):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    _register_fonts()
    rtl = lang == "ar"
    ta = TA_RIGHT if rtl else TA_LEFT
    base = ParagraphStyle("base", fontName=_LATIN, fontSize=9, alignment=ta,
                          leading=12)
    title_style = ParagraphStyle("title", parent=base, fontName=_LATIN_B,
                                 fontSize=16, textColor=colors.HexColor(CDS_RED))
    sub_style = ParagraphStyle("sub", parent=base, fontSize=8,
                               textColor=colors.HexColor(CDS_DARK_GRAY))
    cell = ParagraphStyle("cell", parent=base, fontSize=7.5, leading=9)
    head_cell = ParagraphStyle("hcell", parent=cell, fontName=_LATIN_B,
                               textColor=colors.white)

    cols = COLUMNS[report_key]
    if rtl:
        cols = list(reversed(cols))
    headers = [Paragraph(_pdf_title(ar if rtl else en), head_cell)
               for _k, en, ar, _t, _w in cols]
    data = [headers]
    for row in rows:
        data.append(
            [
                Paragraph(
                    _pdf_text(_cell_value(row, k, t, lang)), cell
                )
                for k, _e, _a, t, _w in cols
            ]
        )

    doc = SimpleDocTemplate(
        _pdf_buf := BytesIO(),
        pagesize=landscape(A4),
        rightMargin=12 * mm, leftMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=14 * mm,
        title=TITLES[report_key][lang],
    )
    story = [
        Paragraph(_pdf_title("CDS Business"), ParagraphStyle(
            "brand", parent=base, fontName=_LATIN_B, fontSize=10,
            textColor=colors.HexColor(CDS_RED))),
        Paragraph(_pdf_title(TITLES[report_key][lang]), title_style),
        Spacer(1, 2 * mm),
        Paragraph(
            _pdf_text(
                f"{'تاريخ الإنشاء' if rtl else 'Generated'}: "
                f"{timezone.localtime().strftime('%Y-%m-%d %H:%M')}"
            ),
            sub_style,
        ),
    ]
    f = meta.get("filters", {})
    rng = " – ".join(x for x in (f.get("date_from"), f.get("date_to")) if x)
    if rng:
        story.append(Paragraph(_pdf_text(
            f"{'الفترة' if rtl else 'Period'}: {rng}"), sub_style))
    story.append(Paragraph(_pdf_text(
        f"{'الفلاتر' if rtl else 'Filters'}: "
        f"{_filters_text(f, lang)}"), sub_style))
    story.append(Spacer(1, 3 * mm))

    pairs = _summary_pairs(report_key, summary, lang)
    if pairs:
        s_rows = [[Paragraph(_pdf_title("Summary" if not rtl else "الملخص"),
                             head_cell), ""]]
        for label, value in pairs:
            s_rows.append([
                Paragraph(_pdf_text(label), cell),
                Paragraph(_pdf_text(value), cell),
            ])
        st = Table(s_rows, colWidths=[60 * mm, 40 * mm],
                   hAlign="RIGHT" if rtl else "LEFT")
        st.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(CDS_RED)),
            ("SPAN", (0, 0), (-1, 0)),
            ("GRID", (0, 1), (-1, -1), 0.25, colors.HexColor("#E7E9EC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white,
                                                  colors.HexColor("#F7F8F9")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story += [st, Spacer(1, 4 * mm)]

    weights = [w for *_r, w in cols]
    usable = doc.width
    total_w = sum(weights)
    tbl = Table(
        data,
        colWidths=[usable * w / total_w for w in weights],
        repeatRows=1,
        hAlign="RIGHT" if rtl else "LEFT",
    )
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(CDS_RED)),
        ("GRID", (0, 1), (-1, -1), 0.25, colors.HexColor("#E7E9EC")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.HexColor(CDS_DEEP_RED)),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#F7F8F9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
    ]))
    story.append(tbl)

    page_lbl = "صفحة" if rtl else "Page"

    def _footer(canv, _doc):
        canv.saveState()
        canv.setFont(_LATIN, 7.5)
        canv.setFillColor(colors.HexColor(CDS_DARK_GRAY))
        txt = f"{page_lbl} {canv.getPageNumber()}"
        canv.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm, txt)
        canv.drawString(12 * mm, 8 * mm, "CDS Business")
        canv.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return _pdf_buf.getvalue()


def export_filename(report_key, fmt, f):
    span = ""
    if f.date_from or f.date_to:
        span = f"_{f.date_from or '…'}_{f.date_to or '…'}"
    return f"report_{report_key}{span}.{fmt}"
